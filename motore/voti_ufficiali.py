"""I voti ufficiali della lega, dai file che Marco scarica da Fantacalcio.it.

Marco e' iscritto e scarica un file per giornata (uno stesso file porta tre
fogli: Fantacalcio, Statistico, Italia — tre modi di dare i voti). Il nostro e'
**Fantacalcio**: e' quello che Leghe usa per calcolare i punti, e il
regolamento della Lega Minuetto dice che il risultato si fa con i voti
pubblicati da Leghe.

Perche' questa e' la fonte migliore che abbiamo:
- e' la stessa che assegna i punti in campionato, non un termometro;
- porta i bonus gia' contati per colonna (gol, assist, rigori, ammonizioni),
  quindi il fantavoto lo calcoliamo esatto invece di dedurlo dal racconto;
- ha il **codice giocatore**, lo stesso del file delle quotazioni: l'abbinamento
  e' per codice, non per cognome, e gli omonimi smettono di essere un problema.

I file restano sul Mac: sono a uso personale degli iscritti, non vanno
pubblicati (`.gitignore` copre gli xlsx). Nel repository pubblico non finisce
nessun voto.

Uso:
    python3 motore/voti_ufficiali.py                    # tutti i file nella cartella
    python3 motore/voti_ufficiali.py "Voti Fantacalcio/Voti_..._Giornata_4.xlsx"
    python3 motore/voti_ufficiali.py --prova            # non scrive niente

Scrive `dati/voti-ufficiali/giornata-N.json` nel formato di dossier.py
(fonte "leghe"), pronto per `dossier.py importa`.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

RADICE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RADICE / "motore"))
from regole import fantavoto  # noqa: E402

LISTONE = RADICE / "report" / "listone.json"
CARTELLA_FILE = RADICE / "Voti Fantacalcio"
USCITA = RADICE / "dati" / "voti-ufficiali"

# Il foglio da leggere. Gli altri due sono scale diverse della stessa giornata:
# se un domani lo staff dicesse che la lega gioca con "Statistico" o "Italia",
# si cambia qui e si rifa' l'importazione, il resto del codice non cambia.
FOGLIO = "Fantacalcio"

COLONNE = ["Cod.", "Ruolo", "Nome", "Voto", "Gf", "Gs", "Rp", "Rs", "Rf", "Au", "Amm", "Esp", "Ass"]

# Come si racconta un evento all'app, in ordine di importanza.
COME_SI_DICE = [("Gf", "gol"), ("Rf", "rigore segnato"), ("Ass", "assist"),
                ("Rp", "rigore parato"), ("Rs", "rigore sbagliato"), ("Au", "autogol"),
                ("Gs", "gol subito"), ("Amm", "ammonito"), ("Esp", "espulso")]


def carica_listone():
    if not LISTONE.exists():
        sys.exit(f"Manca {LISTONE}: prima python3 motore/listone.py")
    giocatori = json.load(open(LISTONE, encoding="utf-8"))["listone"]
    per_codice = {g["id"]: g for g in giocatori if "id" in g}
    if not per_codice:
        sys.exit("Il listone non ha i codici giocatore: rilancia listone.py sul file "
                 "delle quotazioni (la colonna Id) e riprova.")
    return per_codice


def racconta(eventi, ruolo):
    """Gli eventi in parole, per la riga del voto nell'app: '2 gol, ammonito'."""
    pezzi = []
    for chiave, parola in COME_SI_DICE:
        quanti = int(eventi.get(chiave) or 0)
        if not quanti or (chiave == "Gs" and ruolo != "P"):
            continue
        if chiave in ("Amm", "Esp"):
            pezzi.append(parola)
        elif quanti == 1:
            pezzi.append(parola)
        else:
            pezzi.append(f"{quanti} {parola}{'i' if parola.endswith('o') else ''}")
    return ", ".join(pezzi)


def giornata_dal_nome(percorso):
    m = re.search(r"Giornata[_ ]*(\d+)", Path(percorso).name, re.I)
    return int(m.group(1)) if m else None


def leggi_file(percorso, per_codice):
    """Le righe del foglio: una squadra alla volta, poi i suoi giocatori."""
    libro = openpyxl.load_workbook(percorso, data_only=True)
    if FOGLIO not in libro.sheetnames:
        sys.exit(f"{Path(percorso).name}: manca il foglio '{FOGLIO}' (ci sono: {libro.sheetnames})")
    foglio = libro[FOGLIO]
    voti, mancati, squadra = [], [], None
    for riga in foglio.iter_rows(values_only=True):
        prima = riga[0]
        if isinstance(prima, str) and riga[1] is None and len(prima) < 30:
            squadra = prima.strip()           # riga con il nome della squadra
            continue
        if not isinstance(prima, int) or riga[3] in (None, ""):
            continue                          # intestazioni, avvisi, righe vuote
        dati = dict(zip(COLONNE, riga))
        g = per_codice.get(int(prima))
        if not g:
            mancati.append({"codice": int(prima), "nome_file": dati.get("Nome"), "squadra": squadra})
            continue
        ruolo = g["r"]
        # "6*" e' il voto d'ufficio di chi entra per pochi minuti: vale 6 se lo
        # schieri, ma non dice niente su come ha giocato. Lo teniamo (e' una
        # presenza vera) marcandolo, cosi' resta riconoscibile se un domani
        # decidiamo di tenerlo fuori dalle medie.
        grezzo = str(dati["Voto"]).strip()
        ufficio = grezzo.endswith("*")
        voto_num = float(grezzo.rstrip("*").replace(",", "."))
        dati["Voto"] = voto_num
        voce = {"nome": g["n"], "voto": voto_num,
                "fantavoto": fantavoto(voto_num, dati, ruolo)}
        parole = racconta(dati, ruolo)
        if ufficio:
            parole = (parole + ", " if parole else "") + "voto d'ufficio"
        if parole:
            voce["eventi"] = parole
        voti.append(voce)
    return voti, mancati


def scrivi(percorso, per_codice, prova=False):
    giornata = giornata_dal_nome(percorso)
    if not giornata:
        print(f"  {Path(percorso).name}: non capisco di che giornata sia, salto")
        return 0
    voti, mancati = leggi_file(percorso, per_codice)
    for v in voti:
        v["giornata"] = giornata
    dati = {"data": f"giornata-{giornata}", "fonte": "leghe", "note": [], "rigoristi": [],
            "voti": voti, "non_abbinati": mancati,
            "sintesi": f"Voti ufficiali di Fantacalcio.it (foglio {FOGLIO}), {giornata}ª giornata."}
    if not prova:
        USCITA.mkdir(parents=True, exist_ok=True)
        (USCITA / f"giornata-{giornata}.json").write_text(
            json.dumps(dati, ensure_ascii=False, indent=1), encoding="utf-8")
    fuori = f", {len(mancati)} fuori listone" if mancati else ""
    print(f"  giornata {giornata}: {len(voti)} voti{fuori}")
    return len(voti)


def main():
    argomenti = [a for a in sys.argv[1:] if not a.startswith("--")]
    prova = "--prova" in sys.argv
    per_codice = carica_listone()
    file = [Path(a) for a in argomenti] or sorted(CARTELLA_FILE.glob("*.xlsx"))
    if not file:
        sys.exit(f"Nessun file in {CARTELLA_FILE}: scaricali da Fantacalcio.it, uno per giornata.")
    totale = sum(scrivi(f, per_codice, prova) for f in file)
    print(f"{totale} voti ufficiali" + (" (prova, niente scritto)" if prova else
          f" → {USCITA.relative_to(RADICE)}/; ora: python3 motore/dossier.py importa <file>"))


if __name__ == "__main__":
    main()
